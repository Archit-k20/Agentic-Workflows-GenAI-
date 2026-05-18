from io import BytesIO

import docx
import fitz
import streamlit as st
from openai import OpenAI


MAX_FILES = 3
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 180


def split_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    chunks = []
    start = 0
    cleaned = " ".join(text.split())
    while start < len(cleaned):
        end = start + chunk_size
        chunks.append(cleaned[start:end])
        if end >= len(cleaned):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk.strip()]


def extract_pdf_text(uploaded_file):
    uploaded_file.seek(0)
    text_parts = []
    with fitz.open(stream=uploaded_file.read(), filetype="pdf") as pdf:
        for page in pdf:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def extract_docx_text(uploaded_file):
    uploaded_file.seek(0)
    document = docx.Document(BytesIO(uploaded_file.read()))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def uploaded_file_to_documents(uploaded_file):
    from langchain_core.documents import Document

    file_name = uploaded_file.name
    lower_name = file_name.lower()

    if lower_name.endswith(".pdf"):
        text = extract_pdf_text(uploaded_file)
    elif lower_name.endswith(".docx"):
        text = extract_docx_text(uploaded_file)
    else:
        raise ValueError(f"Unsupported file type: {file_name}")

    if not text:
        raise ValueError(f"No readable text found in {file_name}")

    return [
        Document(
            page_content=chunk,
            metadata={"source": file_name, "chunk": index + 1},
        )
        for index, chunk in enumerate(split_text(text))
    ]


def build_document_vectorstore(uploaded_files, api_key):
    from langchain_community.vectorstores import FAISS
    from langchain_openai import OpenAIEmbeddings

    documents = []
    errors = []

    for uploaded_file in uploaded_files[:MAX_FILES]:
        try:
            documents.extend(uploaded_file_to_documents(uploaded_file))
        except Exception as exc:
            errors.append(f"{uploaded_file.name}: {exc}")

    if not documents:
        raise ValueError("No documents could be processed. Please upload readable PDF or DOCX files.")

    embeddings = OpenAIEmbeddings(api_key=api_key)
    vectorstore = FAISS.from_documents(documents, embeddings)
    return vectorstore, errors


def answer_with_context(question, docs, api_key):
    context_blocks = []
    for index, doc in enumerate(docs, start=1):
        source = doc.metadata.get("source", "Uploaded document")
        chunk = doc.metadata.get("chunk", index)
        context_blocks.append(f"[S{index}] Source: {source}, chunk {chunk}\n{doc.page_content}")

    context = "\n\n".join(context_blocks)
    client = OpenAI(api_key=api_key)
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        temperature=0.2,
        messages=[
            {
                "role": "system",
                "content": (
                    "Answer using only the provided document context. "
                    "If the answer is not present, say that the documents do not contain enough information. "
                    "Cite sources inline using labels like [S1]."
                ),
            },
            {
                "role": "user",
                "content": f"Question: {question}\n\nDocument context:\n{context}",
            },
        ],
    )
    return response.choices[0].message.content or "No answer generated."


def show_document_qna(uploaded_files, api_key=None):
    st.title("Document QnA Tool")

    if not api_key:
        st.error("OpenAI API key is required")
        return

    if not uploaded_files:
        st.info("Please upload at least one document to begin.")
        return

    if "document_qna_vectorstore" not in st.session_state:
        st.session_state.document_qna_vectorstore = None
    if "document_qna_errors" not in st.session_state:
        st.session_state.document_qna_errors = []

    if st.button("Process Files"):
        with st.spinner("Processing documents..."):
            try:
                vectorstore, errors = build_document_vectorstore(uploaded_files, api_key)
                st.session_state.document_qna_vectorstore = vectorstore
                st.session_state.document_qna_errors = errors
                st.success("Documents processed and ready for querying.")
            except Exception as exc:
                st.session_state.document_qna_vectorstore = None
                st.error(f"Error processing files: {exc}")
                return

    for error in st.session_state.document_qna_errors:
        st.warning(error)

    query = st.text_input("Ask a question about your documents:")
    if query:
        vectorstore = st.session_state.document_qna_vectorstore
        if vectorstore is None:
            st.warning("Please process documents first.")
            return

        with st.spinner("Searching for answers..."):
            try:
                retrieved_docs = vectorstore.similarity_search(query, k=4)
                answer = answer_with_context(query, retrieved_docs, api_key)
                st.markdown("### Answer")
                st.write(answer)
                st.markdown("### Retrieved Sources")
                for index, doc in enumerate(retrieved_docs, start=1):
                    source = doc.metadata.get("source", "Uploaded document")
                    chunk = doc.metadata.get("chunk", index)
                    st.code(f"[S{index}] {source}, chunk {chunk}")
            except Exception as exc:
                st.error(f"Error answering question: {exc}")
