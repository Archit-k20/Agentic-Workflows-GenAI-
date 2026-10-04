"""Bounded extraction, complete chunk coverage, and local embedding profiles."""

import json
import re
from pathlib import Path
from functools import lru_cache
from . import free_config as cfg
from .events import emit


@lru_cache(maxsize=1)
def tokenizer():
    from transformers import AutoTokenizer

    with cfg.CPU_GATE:
        return AutoTokenizer.from_pretrained(
            cfg.EMBED_MODEL, revision=cfg.EMBED_REVISION
        )


def check_text(text, runtime, source):
    if not text or not text.strip():
        raise ValueError(f"No readable text found in {source}.")
    runtime.check()
    tokens = tokenizer().encode(text, add_special_tokens=False, verbose=False)
    if len(tokens) > cfg.MAX_TOKENS:
        raise ValueError(
            f"{source} exceeds the 12,000-token processing limit. Split the input; nothing was silently truncated."
        )
    runtime.coverage.append({"source": source, "tokens": len(tokens), "complete": True})
    return text


def chunks(text, size=3000, overlap=200):
    # Conservative generation chunks use at most 3,000 UTF-8 bytes, hence at most
    # 3,000 byte-level generation tokens. Original Unicode text is retained.
    if not 0 <= overlap < size:
        raise ValueError("Chunk overlap must be smaller than its size.")
    if size == 384:
        tokens = tokenizer().encode(text, add_special_tokens=False, verbose=False)
        return [
            tokenizer().decode(tokens[start : start + size], skip_special_tokens=True)
            for start in range(0, len(tokens), size - overlap)
        ]
    raw = text.encode("utf-8")
    pieces = []
    start = 0
    while start < len(raw):
        end = min(len(raw), start + size)
        while end < len(raw) and raw[end] & 0xC0 == 0x80:
            end -= 1
        if end == start:
            raise ValueError("The chunk size cannot fit this Unicode character.")
        if end < len(raw):
            # Avoid turning clipped words into invented entities in summaries.
            boundary_start = max(start + overlap + 1, end - 200)
            boundaries = list(re.finditer(rb"[ \t\r\n]", raw[boundary_start:end]))
            if boundaries:
                end = boundary_start + boundaries[-1].end()
        pieces.append(raw[start:end].decode("utf-8"))
        if end == len(raw):
            break
        next_start = max(start + 1, end - overlap)
        if next_start < end and raw[next_start - 1] not in b" \t\r\n":
            boundary = re.search(rb"[ \t\r\n]", raw[next_start:end])
            if boundary:
                next_start += boundary.end()
        start = next_start
        while raw[start] & 0xC0 == 0x80:
            start += 1
    return pieces


def extract(file, runtime):
    file.seek(0)
    suffix = Path(file.name).suffix.lower()
    runtime.expensive = True
    emit("stage", name="Extract complete bounded document", status="running")
    with cfg.CPU_GATE:
        runtime.check()
        if suffix == ".pdf":
            import fitz
            import pytesseract
            from PIL import Image

            with fitz.open(stream=file.read(), filetype="pdf") as pdf:
                if len(pdf) > 50:
                    raise ValueError(
                        "PDF exceeds 50 pages. Split the file before processing."
                    )
                parts = []
                for i, page in enumerate(pdf):
                    runtime.check()
                    text = page.get_text().strip()
                    if not text:
                        emit("stage", name=f"OCR page {i+1}", status="running")
                        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                        text = pytesseract.image_to_string(
                            Image.frombytes(
                                "RGB", [pix.width, pix.height], pix.samples
                            ),
                            timeout=min(
                                30,
                                max(
                                    1, runtime.deadline - __import__("time").monotonic()
                                ),
                            ),
                        )
                        emit("stage", name=f"OCR page {i+1}", status="completed")
                    parts.append(text)
                text = "\n".join(parts)
        elif suffix == ".docx":
            from docx import Document

            doc = Document(file)
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [
                    " | ".join(c.text for c in row.cells)
                    for table in doc.tables
                    for row in table.rows
                ]
            )
        elif suffix == ".txt":
            text = file.read().decode("utf-8", errors="replace")
        else:
            from PIL import Image
            import pytesseract

            text = pytesseract.image_to_string(Image.open(file), timeout=30)
    check_text(text, runtime, file.name)
    emit("stage", name="Extract complete bounded document", status="completed")
    return {"name": file.name, "source_type": suffix.lstrip("."), "text": text}


def article(url, runtime):
    from workflows.url_qna import fetch_article_text

    runtime.expensive = True
    title, text = fetch_article_text(url)
    return {"url": url, "title": title, "text": check_text(text, runtime, url)}


def summarize(text, runtime, source="Input"):
    check_text(text, runtime, source)
    parts = chunks(text)
    summaries = []
    for i, part in enumerate(parts):
        emit("stage", name=f"Summarize section {i+1}/{len(parts)}", status="running")
        summaries.append(
            runtime.text(
                [
                    {
                        "role": "system",
                        "content": "Summarize this source section faithfully. Retain names, dates, numbers, decisions and caveats. Do not add facts.",
                    },
                    {"role": "user", "content": part},
                ],
                limit=512,
            )
        )
        emit("stage", name=f"Summarize section {i+1}/{len(parts)}", status="completed")
    # Repeated bounded reductions cover every section, including long inputs.
    rounds = 0
    while len("\n".join(summaries).encode()) > 6000:
        rounds += 1
        if rounds > 4:
            raise ValueError(
                "Section summaries could not be reduced within the synthesis budget. Split the input and retry; no sections were discarded."
            )
        summaries = [
            runtime.text(
                [
                    {
                        "role": "system",
                        "content": "Combine section summaries faithfully. Preserve names, dates, amounts, actions and caveats.",
                    },
                    {"role": "user", "content": part},
                ],
                limit=512,
            )
            for part in chunks("\n".join(summaries))
        ]
        runtime.check()
    emit("stage", name="Combine complete summary", status="running")
    result = runtime.text(
        [
            {
                "role": "system",
                "content": "Write one clear summary of all supplied sections. Preserve significant dates, names, quantities and caveats; do not invent details.",
            },
            {"role": "user", "content": "\n".join(summaries)},
        ],
        limit=600,
    )
    emit("stage", name="Combine complete summary", status="completed")
    return result


def override(module, name, args, kwargs, fn, runtime):
    """Free hooks reuse legacy reasoning functions with bounded, complete sources."""
    from . import grounding as g
    if name in {"fetch_source", "fetch_support_source"}:
        return article(args[0], runtime)
    if name == "extract_document_text":
        return extract(args[0], runtime)
    if name == "generate_speech":
        try:
            return runtime.speech(
                args[0], kwargs.get("voice", "af_heart"), kwargs["output_path"]
            )
        except Exception as exc:
            runtime.warning("Narration failed; the content artifacts are retained.")
            return "Error: " + str(exc)
    if name == "analyze_document":
        document, client = args
        analyses = []
        for part in chunks(document["text"]):
            with g.scope(runtime, "document", part):
                analysis = fn({**document, "text": part}, client)
            analyses.append(g.document_analysis(analysis, part, runtime))
        if len(analyses) == 1:
            return analyses[0]
        result = {
            "document_type": analyses[0]["document_type"],
            "summary": "\n".join(a["summary"] for a in analyses),
            "entities": {},
            "action_items": [],
            "risks": [],
        }
        for key in ("people", "organizations", "emails", "dates"):
            result["entities"][key] = list(
                dict.fromkeys(x for a in analyses for x in a["entities"][key])
            )
        seen = set()
        for a in analyses:
            for item in a["action_items"]:
                identity = json.dumps(item, sort_keys=True)
                if identity not in seen:
                    result["action_items"].append(item)
                    seen.add(identity)
        result["risks"] = list(dict.fromkeys(x for a in analyses for x in a["risks"]))
        result["summary"] = summarize(
            result["summary"], runtime, "Document section summaries"
        )
        return g.document_analysis(result, document["text"], runtime)
    if name in {"summarize_source", "summarize_support_source"}:
        index = 0 if name == "summarize_source" else 1
        source = args[index]
        results = []
        for part in chunks(source["text"]):
            bounded = list(args)
            bounded[index] = {**source, "text": part}
            with g.scope(runtime, "support-source" if name == "summarize_support_source" else "research-source", part):
                results.append(fn(*bounded, **kwargs))
        result = results[0]
        result["label"] = args[4] if name == "summarize_source" else args[2]
        if len(results) > 1:
            # Support reasoning receives only actual source excerpts, never a
            # second paraphrase that could import a customer's allegation.
            combined = "\n".join(r["summary"] for r in results)
            result["summary"] = combined if name == "summarize_support_source" else summarize(
                combined, runtime, "Source section summaries",
            )
            if name == "summarize_source":
                result["key_points"] = list(
                    dict.fromkeys(x for r in results for x in r.get("key_points", []))
                )
        if name == "summarize_support_source":
            result["metadata"] = {"model_relevance": result.get("relevance", ""),
                                  "source_provenance": "Verbatim documentation excerpt"}
            result["relevance"] = "Quoted documentation; assess eligibility at the resolution stage."
        # Keep draft and review prompts within the local context envelope.
        if len(json.dumps(result).encode()) > 5000:
            runtime.warning(
                "Source digest exceeded the bounded synthesis budget. Split the source before retrying."
            )
            raise ValueError(
                "Source digest exceeds the bounded synthesis budget. Split the source; no key points were silently discarded."
            )
        return result
    if name in {"synthesize_report", "critique_report", "revise_report"}:
        sources = args[2] if name != "critique_report" else args[1]
        with g.scope(runtime, "research-report", labels=[s["label"] for s in sources]):
            result = fn(*args, **kwargs)
        if name == "critique_report":
            return g.research_review(result, args[2], sources, runtime)
        issues = g.report_issues(result, sources)
        if issues:
            if name == "revise_report" and not g.report_issues(args[3], sources):
                runtime.warning("Revision failed report structure/citation checks; the last valid report was retained.")
                history = getattr(runtime, "research_rejections", [])
                history.append({"stage": "revision", "revision": result, "issues": issues})
                runtime.research_rejections = history
                return args[3]
            raise ValueError("Research report failed source/structure checks: " + "; ".join(issues))
        return result
    if name == "draft_support_resolution":
        question, intent, sources, _ = args
        source = "\n".join(s["summary"] for s in sources)
        with g.scope(runtime, "support-draft", source, [s["label"] for s in sources]):
            bounded = list(args)
            bounded[2] = [{k: v for k, v in item.items() if k != "metadata"} for item in sources]
            result = fn(*bounded, **kwargs)
        return g.support_draft(result, question, intent, sources, runtime)
    if name == "finalize_triage":
        result = fn(*args, **kwargs)
        return g.support_final(result, args[2], args[3], runtime)
    if name in {"build_content_plan", "build_content_package", "critique_content_package"}:
        with g.scope(runtime, "content", args[0]):
            result = fn(*args, **kwargs)
        if name == "build_content_plan":
            result["proposal_scope"] = "Creative proposals only; the original idea is the factual source."
        elif name == "build_content_package":
            output = "\n".join([result["title"], result["script"], result["cta"], *result["captions"].values(), *result["hashtags"]])
            issues = g.claim_issues(output, args[0])
            if issues:
                raise ValueError("Content package failed source checks: " + "; ".join(issues))
            result["script"], _ = g.restore_details(result["script"], args[0], runtime, "Content script")
        else:
            return g.content_review(result, args[0], args[4], runtime)
        return result
    return fn(*args, **kwargs)
